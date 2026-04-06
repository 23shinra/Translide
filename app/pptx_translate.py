from __future__ import annotations

from io import BytesIO
from typing import Optional

from .translate import TranslateClient


def _walk_shapes(shape):
    """Группы, таблицы и обычные фигуры — рекурсивно."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            yield from _walk_shapes(child)
        return

    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                tf = getattr(cell, "text_frame", None)
                if tf is not None:
                    yield tf
        return

    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        yield shape.text_frame


def _iter_text_frames(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            yield from _walk_shapes(shape)
        if getattr(slide, "has_notes_slide", False) and slide.has_notes_slide:
            notes = slide.notes_slide
            ntf = getattr(notes, "notes_text_frame", None)
            if ntf is not None:
                yield ntf


def _runs_with_text(paragraph) -> list:
    return [r for r in paragraph.runs if (r.text or "").strip()]


async def _translate_paragraph(
    paragraph,
    *,
    translator: TranslateClient,
    source_lang: Optional[str],
    target_lang: str,
) -> None:
    """
    Если в абзаце один непрерывный фрагмент текста — переводим целиком (как в прежней версии).
    Если текст разбит на несколько run (часто так в PowerPoint) — переводим каждый run отдельно,
    чтобы не склеивались слова без пробелов и сохранялось локальное форматирование.
    """
    runs = list(paragraph.runs)
    with_text = _runs_with_text(paragraph)
    if not with_text:
        return

    if len(with_text) <= 1:
        original = paragraph.text or ""
        if not original.strip():
            return
        translated = await translator.translate(
            original,
            source_lang=source_lang,
            target_lang=target_lang,
        )
        if runs:
            runs[0].text = translated
            for r in runs[1:]:
                r.text = ""
        else:
            paragraph.text = translated
        return

    for r in runs:
        chunk = r.text or ""
        if not chunk.strip():
            continue
        r.text = await translator.translate(
            chunk,
            source_lang=source_lang,
            target_lang=target_lang,
        )


async def translate_pptx_bytes(
    *,
    pptx_bytes: bytes,
    source_lang: Optional[str],
    target_lang: str,
    translator: TranslateClient,
) -> bytes:
    from pptx import Presentation  # lazy import

    prs = Presentation(BytesIO(pptx_bytes))

    for text_frame in _iter_text_frames(prs):
        for paragraph in text_frame.paragraphs:
            await _translate_paragraph(
                paragraph,
                translator=translator,
                source_lang=source_lang,
                target_lang=target_lang,
            )

    out = BytesIO()
    prs.save(out)
    return out.getvalue()
